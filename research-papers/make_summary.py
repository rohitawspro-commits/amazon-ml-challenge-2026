"""Build one combined IEEE-format summary of both NLP papers (single .docx)."""
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from make_papers import AUTHOR, PAPER1, PAPER2, bullet, heading, para, set_cols, subheading

# References shared by both papers: Paper 2 number -> Paper 1 number.
SAME = {8: 6, 9: 7, 10: 9}
REFS = {f"P1-{i}": r for i, r in enumerate(PAPER1["refs"], 1)}
REFS.update({f"P2-{i}": r for i, r in enumerate(PAPER2["refs"], 1) if i not in SAME})

order = []  # citation keys in order of first appearance


def key(paper, n):
    return f"P1-{SAME[n]}" if paper == 2 and n in SAME else f"P{paper}-{n}"


def tag(text, paper):
    """Rewrite a paper's local [n] / [a]-[b] citations as global keys."""
    text = re.sub(r"\[(\d+)\]–\[(\d+)\]",
                  lambda m: "[" + ",".join(key(paper, n) for n in range(int(m[1]), int(m[2]) + 1)) + "]",
                  text)
    return re.sub(r"\[(\d+)\]", lambda m: f"[{key(paper, int(m[1]))}]", text)


def cite(text):
    """Replace global keys with final numbers, assigned by first appearance."""
    def repl(m):
        nums = []
        for k in m[1].split(","):
            if k not in order:
                order.append(k)
            nums.append(order.index(k) + 1)
        nums.sort()
        out, i = [], 0
        while i < len(nums):
            j = i
            while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
                j += 1
            out.append(f"[{nums[i]}]–[{nums[j]}]" if j - i >= 2 else
                       ", ".join(f"[{n}]" for n in nums[i:j + 1]))
            i = j + 1
        return ", ".join(out)
    return re.sub(r"\[((?:P[12]-\d+)(?:,P[12]-\d+)*)\]", repl, text)


SECTION_MAP = {
    "I. Introduction": "A. Problem Statement and Objectives",
    "II. Literature Survey": "B. Literature Survey",
    "III. Proposed Methodology": "C. Proposed Methodology",
}
SKIP = ("The rest of the paper", "Section II presents")


def paper_items(paper, n):
    items, step = [], 0
    for kind, text in paper["body"]:
        if kind == "p" and text.startswith(SKIP):
            continue
        if kind == "h":
            items.append(("sh", SECTION_MAP[text]))
            step = 0
        elif kind == "sh":
            step += 1
            items.append(("lead", f"{step}) {text.split('. ', 1)[1]}: "))
        else:
            items.append((kind, tag(text, n)))
    return items


INTRO = [
    ("h", "I. Introduction"),
    ("p", "Natural Language Processing (NLP) is the branch of Artificial Intelligence that "
          "enables computers to read, understand and generate human language. Every day, huge "
          "amounts of text are produced on e-commerce websites, news portals and social media, "
          "and NLP techniques are needed to turn this unstructured text into useful information. "
          "Text classification is one of the most important NLP tasks, where a piece of text is "
          "assigned to one of several predefined categories."),
    ("p", "This document summarises two research papers that both apply text classification to "
          "real-world problems. The first paper, “Sentiment Analysis of Customer Product "
          "Reviews Using Transformer-Based Language Models”, classifies online product reviews "
          "as positive, neutral or negative. The second paper, “Fake News Detection Using "
          "Natural Language Processing and Deep Learning Techniques”, classifies news "
          "articles as real or fake. Both papers compare classical machine learning, recurrent "
          "and convolutional deep learning, and transformer-based models such as BERT [P1-9]."),
    ("p", "Section II summarises the first paper and Section III summarises the second paper. "
          "Section IV compares the two works, Section V explains the NLP techniques common to "
          "both, Section VI discusses challenges and future scope, and Section VII concludes "
          "the summary."),
]

P1_HEAD = [("h", "II. Paper 1: Sentiment Analysis of Customer Product Reviews")]
P1_TAIL = [
    ("sh", "D. Key Contributions"),
    ("b", "A complete pipeline for noisy e-commerce review text, including negation-aware "
          "preprocessing."),
    ("b", "Fine-tuning of BERT for three-class review sentiment classification."),
    ("b", "A fair comparison of TF-IDF, Bi-LSTM and transformer models on the same data split."),
]
P2_HEAD = [("h", "III. Paper 2: Fake News Detection Using NLP and Deep Learning")]
P2_TAIL = [
    ("sh", "D. Key Contributions"),
    ("b", "A content-based fake news classifier evaluated on both long articles (ISOT) and "
          "short statements (LIAR)."),
    ("b", "Removal of publisher cues so that the model learns from language, not from the "
          "source name."),
    ("b", "Comparison of classical, hybrid CNN–LSTM and BERT models with special focus on "
          "recall of the fake class."),
]

TABLE = [
    ("Aspect", "Paper 1: Sentiment Analysis", "Paper 2: Fake News Detection"),
    ("Task", "Multi-class (positive / neutral / negative)", "Binary (real / fake)"),
    ("Input text", "Short product reviews", "News headline and article body"),
    ("Datasets", "Amazon product reviews", "ISOT and LIAR"),
    ("Classical models", "Naive Bayes, Logistic Regression", "LR, Naive Bayes, Linear SVM, "
                                                              "Random Forest"),
    ("Deep learning model", "Bi-LSTM with GloVe", "Hybrid CNN–LSTM with GloVe"),
    ("Transformer", "BERT-base, 256 tokens, 3-class head", "BERT-base, 512 tokens, binary head"),
    ("Main metric", "Macro F1-score", "F1-score and recall of fake class"),
    ("Main challenge", "Neutral class, negation, sarcasm", "Topic/source bias, evolving writing "
                                                           "style"),
    ("Application", "Summarising customer feedback", "Supporting fact-checkers"),
]

AFTER_TABLE = [
    ("p", "Table I shows that the two papers follow the same overall research design: a clean "
          "preprocessing pipeline, simple TF-IDF baselines, a neural baseline using pre-trained "
          "word embeddings and finally a fine-tuned transformer. The main difference lies in the "
          "nature of the text. Reviews are short and opinion-heavy, so capturing negation and "
          "intensity is important. News articles are long and factual in style, so the model "
          "must capture writing style and deception cues across many sentences, which is why "
          "Paper 2 uses a longer input length and a CNN–LSTM hybrid."),
    ("p", "Another difference is the cost of errors. In sentiment analysis, confusing neutral "
          "and positive reviews has limited impact, so macro F1-score is a suitable summary "
          "metric. In fake news detection, a fake article that is predicted as real can mislead "
          "many readers [P2-2], so recall of the fake class is treated as a priority."),
    ("h", "V. Common NLP Techniques Used"),
    ("sh", "A. Text Preprocessing"),
    ("p", "Both works clean the raw text by removing HTML tags, URLs and special characters, "
          "then tokenise it. Stop-word removal and lemmatisation are applied only for classical "
          "models, because transformer models need the complete sentence to understand context."),
    ("sh", "B. TF-IDF Features"),
    ("p", "Term Frequency–Inverse Document Frequency [P2-12] gives a high weight to words "
          "that are frequent in a document but rare in the whole collection. Unigram and bigram "
          "TF-IDF vectors with linear classifiers form fast and strong baselines [P1-2], [P2-5]."),
    ("sh", "C. Word Embeddings"),
    ("p", "Word2Vec [P1-4] and GloVe [P1-5] map every word to a dense vector so that words with "
          "similar meaning have similar vectors. These vectors are used as the first layer of "
          "the LSTM and CNN models in both papers."),
    ("sh", "D. CNN and LSTM Networks"),
    ("p", "A CNN [P1-6] slides filters over the word sequence to detect important phrases, while "
          "an LSTM [P1-7] reads the sequence step by step and remembers long-range information "
          "using gates. Paper 1 uses a Bi-LSTM, and Paper 2 combines both into a CNN–LSTM."),
    ("sh", "E. Transformers and BERT"),
    ("p", "The transformer [P1-8] uses self-attention so that every word can directly attend to "
          "every other word. BERT [P1-9] is pre-trained on large corpora with masked language "
          "modelling and is then fine-tuned with a small classification head. Both papers use "
          "the [CLS] token output of bert-base-uncased for classification."),
    ("h", "VI. Challenges and Future Scope"),
    ("p", "The following challenges are common to both works:"),
    ("b", "Sarcasm, irony and code-mixed (Hindi–English) text are difficult for current "
          "models."),
    ("b", "Transformer models need GPUs and are slow on large volumes; smaller models such as "
          "DistilBERT [P1-11] can reduce this cost."),
    ("b", "Models may learn shortcuts from dataset bias, such as publisher names or product "
          "categories, instead of real language patterns."),
    ("b", "Fake news styles change over time, so the detection model must be retrained "
          "regularly."),
    ("p", "Future work includes aspect-based sentiment analysis, which finds opinions about "
          "specific product features such as battery or camera; adding social-context and "
          "user-behaviour features for fake news detection [P2-6], [P2-7]; multilingual models "
          "for Indian languages; and explainability methods that highlight the words responsible "
          "for each prediction."),
    ("h", "VII. Conclusion"),
    ("p", "This document summarised two NLP research papers on sentiment analysis of product "
          "reviews and on fake news detection. Although the applications are different, both "
          "papers treat the problem as text classification and follow a common path from "
          "classical TF-IDF models to deep learning and finally to fine-tuned transformer models "
          "such as BERT. The study shows that pre-trained language models are the most promising "
          "approach for understanding context, negation and writing style, while classical "
          "models remain useful as fast and interpretable baselines. Together, these papers "
          "show how NLP can help both businesses and society make better use of online text."),
]


def run_para(doc, text, lead=None):
    p = para(doc)
    if lead:
        r = p.add_run(lead)
        r.font.size, r.italic = Pt(10), True
    r = p.add_run(cite(text))
    r.font.size = Pt(10)


def render(doc, items):
    lead = None
    for kind, text in items:
        if kind == "h":
            heading(doc, text)
        elif kind == "sh":
            subheading(doc, text)
        elif kind == "lead":
            lead = text
        elif kind == "b":
            bullet(doc, cite(text))
        else:
            run_para(doc, text, lead)
            lead = None


def table(doc):
    para(doc, "Table I", size=8, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, indent=False,
         small_caps=True)
    para(doc, "Comparison of the Two Research Papers", size=8, align=WD_ALIGN_PARAGRAPH.CENTER,
         after=3, indent=False, small_caps=True)
    t = doc.add_table(rows=len(TABLE), cols=3)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = (Inches(0.75), Inches(1.25), Inches(1.25))
    for j, w in enumerate(widths):
        t.columns[j].width = w
    for i, row in enumerate(TABLE):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.width = widths[j]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.size, r.bold = Pt(7.5), i == 0 or j == 0
    para(doc, "", after=4, indent=False)


def build(out):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.font.size = Pt(10)
    s = doc.sections[0]
    s.page_height, s.page_width = Inches(11.69), Inches(8.27)
    s.top_margin, s.bottom_margin = Inches(0.75), Inches(1.0)
    s.left_margin = s.right_margin = Inches(0.625)
    set_cols(s, 1)

    para(doc, "A Summary of Two NLP Research Papers: Sentiment Analysis of Product Reviews "
              "and Fake News Detection", size=22, align=WD_ALIGN_PARAGRAPH.CENTER, after=12,
         indent=False)
    para(doc, AUTHOR["name"], size=11, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    for ln in AUTHOR["lines"]:
        para(doc, ln, size=10, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    para(doc, "", after=6, indent=False)

    set_cols(doc.add_section(WD_SECTION.CONTINUOUS), 2, 360)

    for lead, body in (
        ("Abstract—",
         "This document presents a combined summary of two research papers in the field of "
         "Natural Language Processing (NLP). The first paper proposes a sentiment analysis "
         "system that classifies e-commerce product reviews as positive, neutral or negative by "
         "fine-tuning the BERT transformer model and comparing it with TF-IDF and Bi-LSTM "
         "baselines. The second paper proposes a fake news detection system that classifies "
         "news articles as real or fake using classical machine learning, a hybrid CNN–LSTM "
         "network and a fine-tuned BERT model on the ISOT and LIAR datasets. For each paper, the "
         "problem statement, literature survey, methodology and key contributions are "
         "summarised. The two works are then compared, the NLP techniques common to both are "
         "explained, and the main challenges and future directions are discussed. The summary "
         "shows that both problems can be solved as text classification tasks and that "
         "pre-trained transformer models are the most promising approach for both."),
        ("Keywords—",
         "Natural Language Processing, Text Classification, Sentiment Analysis, Fake News "
         "Detection, BERT, LSTM, TF-IDF, Deep Learning"),
    ):
        p = para(doc)
        r = p.add_run(lead)
        r.font.size, r.bold, r.italic = Pt(9), True, True
        r = p.add_run(body)
        r.font.size, r.bold = Pt(9), True

    render(doc, INTRO)
    render(doc, P1_HEAD + paper_items(PAPER1, 1) + P1_TAIL)
    render(doc, P2_HEAD + paper_items(PAPER2, 2) + P2_TAIL)
    heading(doc, "IV. Comparative Analysis")
    table(doc)
    render(doc, AFTER_TABLE)

    heading(doc, "References")
    for i, k in enumerate(order, 1):
        p = para(doc, indent=False)
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        r = p.add_run(f"[{i}]\t{REFS[k]}")
        r.font.size = Pt(8)
    unused = set(REFS) - set(order)
    assert not unused, f"uncited references: {unused}"
    doc.save(out)


if __name__ == "__main__":
    build("NLP_Research_Papers_Summary_Rohit_Pujari.docx")
