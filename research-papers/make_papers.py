"""Generate two IEEE-format NLP research papers (front two pages) as .docx."""
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

AUTHOR = {
    "name": "Rohit Pujari",
    "lines": [
        "Roll No. 28, B.E. (Final Year)",
        "Department of Computer Engineering",
        "[College Name], [City], India",
        "Guide: Prof. [Guide Name]",
    ],
}


def set_cols(section, num, space_twips=360):
    sect_pr = section._sectPr
    for c in sect_pr.findall(qn("w:cols")):
        sect_pr.remove(c)
    cols = OxmlElement("w:cols")
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), str(space_twips))
    sect_pr.append(cols)


def para(doc, text="", size=10, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
         before=0, after=0, indent=True, small_caps=False):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = 1.0
    if indent:
        pf.first_line_indent = Inches(0.15)
    if text:
        r = p.add_run(text)
        r.font.size, r.bold, r.italic, r.font.small_caps = Pt(size), bold, italic, small_caps
    return p


def lead_para(doc, lead, body, lead_bold=True, lead_italic=True):
    p = para(doc, indent=True)
    r = p.add_run(lead)
    r.font.size, r.bold, r.italic = Pt(9), lead_bold, lead_italic
    r = p.add_run(body)
    r.font.size, r.bold = Pt(9), True
    return p


def heading(doc, text):
    para(doc, text, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, before=8, after=4,
         indent=False, small_caps=True)


def subheading(doc, text):
    para(doc, text, size=10, italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=4, after=2,
         indent=False)


def bullet(doc, text):
    p = para(doc, indent=False)
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    r = p.add_run("• " + text)
    r.font.size = Pt(10)


def build(paper, out):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.font.size = Pt(10)

    s = doc.sections[0]
    s.page_height, s.page_width = Inches(11.69), Inches(8.27)  # A4
    s.top_margin, s.bottom_margin = Inches(0.75), Inches(1.0)
    s.left_margin = s.right_margin = Inches(0.625)
    set_cols(s, 1)

    # Title + author block (single column)
    para(doc, paper["title"], size=24, align=WD_ALIGN_PARAGRAPH.CENTER, after=12, indent=False)
    para(doc, AUTHOR["name"], size=11, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    for ln in AUTHOR["lines"]:
        para(doc, ln, size=10, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    para(doc, "", after=6, indent=False)

    body = doc.add_section(WD_SECTION.CONTINUOUS)
    set_cols(body, 2, 360)

    lead_para(doc, "Abstract—", paper["abstract"])
    lead_para(doc, "Keywords—", paper["keywords"])

    for kind, val in paper["body"]:
        {"h": heading, "sh": subheading, "b": bullet}.get(
            kind, lambda d, t: para(d, t))(doc, val)

    heading(doc, "References")
    for i, ref in enumerate(paper["refs"], 1):
        p = para(doc, indent=False)
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        r = p.add_run(f"[{i}]\t{ref}")
        r.font.size = Pt(8)
    doc.save(out)


PAPER1 = {
    "title": "Sentiment Analysis of Customer Product Reviews Using Transformer-Based Language Models",
    "abstract": (
        "Online shopping platforms receive millions of customer reviews every day, and manually "
        "reading them to understand customer opinion is not practical. Sentiment analysis, a core "
        "task of Natural Language Processing (NLP), automatically identifies whether a piece of text "
        "expresses a positive, negative or neutral opinion. Traditional approaches based on "
        "Bag-of-Words features and classical machine learning classifiers fail to capture context, "
        "negation and sarcasm present in real reviews. This paper proposes a sentiment "
        "classification system for e-commerce product reviews that fine-tunes a pre-trained "
        "transformer model, BERT, and compares it with classical baselines (TF-IDF with Logistic "
        "Regression and Naive Bayes) and a recurrent baseline (Bi-LSTM with GloVe embeddings). The "
        "paper describes the dataset, the text preprocessing pipeline, the model architecture, the "
        "training strategy and the evaluation metrics (accuracy, precision, recall and F1-score). "
        "The proposed system is intended to help sellers and platforms summarise customer "
        "feedback quickly and take data-driven product decisions."
    ),
    "keywords": "Natural Language Processing, Sentiment Analysis, BERT, Transformers, "
                "Deep Learning, Product Reviews, Text Classification",
    "body": [
        ("h", "I. Introduction"),
        ("p", "The rapid growth of e-commerce has changed the way people buy products. Before "
              "purchasing, most customers read the reviews written by earlier buyers, and sellers "
              "depend on the same reviews to understand what customers like or dislike about a "
              "product. Popular products can collect thousands of reviews, which makes manual "
              "analysis slow, costly and inconsistent. Sentiment analysis, also known as opinion "
              "mining, is the branch of Natural Language Processing (NLP) that automatically "
              "extracts the polarity of an opinion from text [1]."),
        ("p", "Early work treated sentiment analysis as a standard text classification problem and "
              "applied Naive Bayes, Maximum Entropy and Support Vector Machine classifiers on "
              "word-count features [2]. These methods are fast and simple but ignore word order, "
              "so sentences such as “not bad at all” or “I expected much better” "
              "are often misclassified. Deep learning models such as Convolutional Neural Networks "
              "(CNN) and Long Short-Term Memory (LSTM) networks improved results by learning "
              "features from word embeddings [4]–[7]. More recently, transformer-based "
              "language models pre-trained on large text corpora, such as BERT [9], have become "
              "the standard approach for most NLP tasks because they model the full context of "
              "every word in both directions."),
        ("p", "The main objectives of this work are:"),
        ("b", "To build a clean, reproducible preprocessing pipeline for noisy product review text."),
        ("b", "To fine-tune a pre-trained BERT model for three-class (positive, neutral, negative) "
              "review sentiment classification."),
        ("b", "To compare the transformer model with classical machine learning and Bi-LSTM "
              "baselines using standard evaluation metrics."),
        ("b", "To design a simple interface through which a user can enter a review and obtain its "
              "predicted sentiment with a confidence score."),
        ("p", "The rest of the paper is organised as follows. Section II reviews related work, "
              "Section III describes the proposed methodology and Section IV explains the system "
              "architecture and evaluation plan."),
        ("h", "II. Literature Survey"),
        ("p", "Pang et al. [2] were among the first to apply machine learning to sentiment "
              "classification of movie reviews and showed that unigram features with SVM performed "
              "better than hand-crafted rules. Maas et al. [3] released the widely used IMDb "
              "Large Movie Review dataset and learned word vectors that capture sentiment "
              "information. Word embedding techniques such as Word2Vec [4] and GloVe [5] "
              "represent words as dense vectors so that semantically similar words lie close to "
              "each other, which became the input layer for most neural models."),
        ("p", "Kim [6] showed that a simple CNN with one convolution layer over pre-trained word "
              "vectors achieves strong results on several sentence classification benchmarks. "
              "LSTM networks [7] handle long-range dependencies in text and have been widely used "
              "for review sentiment analysis. Socher et al. [13] introduced the Stanford Sentiment "
              "Treebank and demonstrated that modelling sentence structure helps to capture "
              "negation."),
        ("p", "The transformer architecture proposed by Vaswani et al. [8] replaced recurrence "
              "with self-attention and allowed efficient training on very large corpora. Devlin "
              "et al. [9] introduced BERT, which is pre-trained using masked language modelling "
              "and can be fine-tuned for downstream tasks by adding a single output layer. "
              "RoBERTa [10] improved BERT through better pre-training choices, while DistilBERT "
              "[11] reduced model size for faster inference. McAuley and Leskovec [12] collected "
              "a large Amazon product review corpus that is commonly used for review-level "
              "research. Based on this survey, fine-tuned transformer models are selected as the "
              "core of the proposed system, with classical and recurrent models as baselines."),
        ("h", "III. Proposed Methodology"),
        ("sh", "A. Dataset"),
        ("p", "The system uses publicly available Amazon product reviews [12]. Each review "
              "contains the review text and a star rating from 1 to 5. Ratings of 1–2 are "
              "labelled negative, 3 is labelled neutral and 4–5 are labelled positive. A "
              "balanced subset is sampled so that each class is equally represented, and the data "
              "is split into training (80%), validation (10%) and test (10%) sets."),
        ("sh", "B. Text Preprocessing"),
        ("p", "Raw review text contains HTML tags, URLs, emojis, repeated characters and spelling "
              "mistakes. The preprocessing steps are: (i) removal of HTML tags and URLs, "
              "(ii) lower-casing (for baseline models), (iii) expansion of contractions such as "
              "“don’t” to “do not” so that negation is preserved, "
              "(iv) tokenisation, and (v) stop-word removal and lemmatisation for the classical "
              "baselines only. For BERT, the WordPiece tokenizer is applied directly and stop "
              "words are kept, because the model relies on the full sentence context."),
        ("sh", "C. Baseline Models"),
        ("p", "Two classical baselines are trained on TF-IDF features of unigrams and bigrams: "
              "Multinomial Naive Bayes and Logistic Regression. A deep learning baseline uses a "
              "two-layer Bidirectional LSTM with 300-dimensional GloVe [5] embeddings followed by "
              "a dense softmax layer."),
        ("sh", "D. Transformer Model"),
        ("p", "The main model is the pre-trained bert-base-uncased model [9], which has 12 "
              "encoder layers, 768 hidden units and 12 attention heads. Each review is converted "
              "to the form [CLS] tokens [SEP], truncated or padded to 256 tokens. The final "
              "hidden state of the [CLS] token is passed through a dropout layer and a linear "
              "layer with softmax activation to produce probabilities for the three classes. The "
              "complete network is fine-tuned end-to-end using the AdamW optimiser with a small "
              "learning rate, linear warm-up and cross-entropy loss. Early stopping on the "
              "validation F1-score is used to prevent over-fitting."),
        ("sh", "E. Evaluation Metrics"),
        ("p", "All models are evaluated on the same held-out test set using accuracy, "
              "precision, recall and macro-averaged F1-score. A confusion matrix is used to study "
              "the errors between neighbouring classes, especially between neutral and the two "
              "polar classes, which is known to be the most difficult case."),
    ],
    "refs": [
        "B. Liu, Sentiment Analysis and Opinion Mining. San Rafael, CA, USA: Morgan & Claypool, 2012.",
        "B. Pang, L. Lee, and S. Vaithyanathan, “Thumbs up? Sentiment classification using "
        "machine learning techniques,” in Proc. EMNLP, 2002, pp. 79–86.",
        "A. L. Maas, R. E. Daly, P. T. Pham, D. Huang, A. Y. Ng, and C. Potts, “Learning word "
        "vectors for sentiment analysis,” in Proc. ACL-HLT, 2011, pp. 142–150.",
        "T. Mikolov, K. Chen, G. Corrado, and J. Dean, “Efficient estimation of word "
        "representations in vector space,” arXiv:1301.3781, 2013.",
        "J. Pennington, R. Socher, and C. D. Manning, “GloVe: Global vectors for word "
        "representation,” in Proc. EMNLP, 2014, pp. 1532–1543.",
        "Y. Kim, “Convolutional neural networks for sentence classification,” in Proc. "
        "EMNLP, 2014, pp. 1746–1751.",
        "S. Hochreiter and J. Schmidhuber, “Long short-term memory,” Neural Computation, "
        "vol. 9, no. 8, pp. 1735–1780, 1997.",
        "A. Vaswani et al., “Attention is all you need,” in Proc. NeurIPS, 2017, "
        "pp. 5998–6008.",
        "J. Devlin, M.-W. Chang, K. Lee, and K. Toutanova, “BERT: Pre-training of deep "
        "bidirectional transformers for language understanding,” in Proc. NAACL-HLT, 2019, "
        "pp. 4171–4186.",
        "Y. Liu et al., “RoBERTa: A robustly optimized BERT pretraining approach,” "
        "arXiv:1907.11692, 2019.",
        "V. Sanh, L. Debut, J. Chaumond, and T. Wolf, “DistilBERT, a distilled version of "
        "BERT: smaller, faster, cheaper and lighter,” arXiv:1910.01108, 2019.",
        "J. McAuley and J. Leskovec, “Hidden factors and hidden topics: Understanding rating "
        "dimensions with review text,” in Proc. ACM RecSys, 2013, pp. 165–172.",
        "R. Socher et al., “Recursive deep models for semantic compositionality over a "
        "sentiment treebank,” in Proc. EMNLP, 2013, pp. 1631–1642.",
    ],
}

PAPER2 = {
    "title": "Fake News Detection Using Natural Language Processing and Deep Learning Techniques",
    "abstract": (
        "Social media and online news portals allow information to spread to millions of users "
        "within minutes, but they also allow false and misleading news to spread just as fast. "
        "Fake news can influence public opinion, elections, health decisions and financial "
        "markets, and manual fact-checking cannot keep up with the volume of content being "
        "published. This paper presents an automatic fake news detection system based on Natural "
        "Language Processing (NLP). The system analyses the headline and body text of a news "
        "article and classifies it as real or fake. Three families of models are studied: "
        "classical machine learning classifiers on TF-IDF features, a hybrid CNN–LSTM deep "
        "learning model using pre-trained word embeddings, and a fine-tuned BERT transformer "
        "model. The paper describes publicly available benchmark datasets, the preprocessing "
        "pipeline, the architecture of each model and the evaluation strategy. The goal of the "
        "work is to build a reliable, explainable tool that can assist readers and fact-checkers "
        "in identifying suspicious news content at an early stage."
    ),
    "keywords": "Fake News Detection, Natural Language Processing, Misinformation, TF-IDF, "
                "CNN-LSTM, BERT, Text Classification",
    "body": [
        ("h", "I. Introduction"),
        ("p", "News consumption has moved from newspapers and television to social media "
              "platforms and messaging applications. While this makes information easily "
              "accessible, it also makes it easy for anyone to publish content without editorial "
              "checks. Fake news is defined as news articles that are intentionally and verifiably "
              "false and could mislead readers [1]. A large-scale study of Twitter by Vosoughi et "
              "al. [2] found that false news spreads farther, faster and deeper than true news, "
              "which shows how serious the problem has become."),
        ("p", "Professional fact-checking organisations verify claims manually, but this process "
              "is slow and does not scale to the huge amount of content published every day. "
              "Therefore, automatic detection methods are required. Since the content of a news "
              "article is written in natural language, NLP techniques can be used to learn the "
              "linguistic patterns, writing style and word usage that differ between fake and "
              "genuine news [4]. Fake articles often use emotional words, exaggerated headlines, "
              "fewer sources and a sensational tone, and these signals can be captured by "
              "machine learning and deep learning models."),
        ("p", "The main objectives of this work are:"),
        ("b", "To study existing approaches and benchmark datasets for fake news detection."),
        ("b", "To design an NLP preprocessing and feature extraction pipeline for news articles."),
        ("b", "To implement and compare classical machine learning, CNN–LSTM and BERT-based "
              "classifiers for real/fake news classification."),
        ("b", "To provide a simple web interface in which a user can paste a news article and "
              "receive a prediction with a confidence score."),
        ("p", "Section II presents the literature survey, Section III explains the datasets and "
              "the proposed methodology, and Section IV describes the system architecture and "
              "evaluation plan."),
        ("h", "II. Literature Survey"),
        ("p", "Shu et al. [1] presented a detailed survey of fake news detection from a data "
              "mining perspective and divided the approaches into news-content-based and "
              "social-context-based methods. Wang [3] released the LIAR dataset containing about "
              "12.8K short political statements labelled with six degrees of truthfulness, and "
              "showed that a hybrid CNN using text and metadata performs better than text-only "
              "models. Pérez-Rosas et al. [4] built datasets covering several news domains "
              "and studied linguistic features such as n-grams, punctuation, readability and "
              "psycholinguistic categories for fake news detection."),
        ("p", "Ahmed et al. [5] used n-gram features with TF-IDF weighting and compared six "
              "classical classifiers, finding that linear models such as Linear SVM work well; "
              "their ISOT dataset is widely used in student and research projects. Shu et al. [6] "
              "introduced FakeNewsNet, which includes news content along with social context from "
              "PolitiFact and GossipCop. Ruchansky et al. [7] proposed CSI, a hybrid deep model "
              "that combines the text of articles, the response of users and the behaviour of "
              "sources."),
        ("p", "Deep learning models remove the need for manual feature engineering. CNNs [8] "
              "capture local n-gram patterns and LSTMs [9] capture long-range dependencies in "
              "text. Transformer-based models such as BERT [10] learn deep contextual "
              "representations and have achieved state-of-the-art results on many NLP tasks. "
              "Kaliyar et al. [11] proposed FakeBERT, which combines BERT embeddings with parallel "
              "one-dimensional convolution blocks for fake news classification. Based on this "
              "survey, the proposed work compares content-based classical, hybrid deep learning "
              "and transformer models on the same datasets."),
        ("h", "III. Proposed Methodology"),
        ("sh", "A. Datasets"),
        ("p", "Two public datasets are used: the ISOT Fake News dataset [5], which contains full "
              "articles labelled as real or fake, and the LIAR dataset [3], which contains short "
              "statements. For LIAR, the six labels are mapped into two classes (true, mostly-true "
              "and half-true as real; barely-true, false and pants-fire as fake). Each dataset is "
              "split into 80% training, 10% validation and 10% test data using stratified sampling."),
        ("sh", "B. Preprocessing and Feature Extraction"),
        ("p", "The headline and body are concatenated to form the input text. The preprocessing "
              "steps include removal of URLs, HTML tags, special characters and extra spaces, "
              "lower-casing, tokenisation, stop-word removal and lemmatisation. Source names and "
              "agency tags such as “(Reuters)” are removed so that the model cannot learn "
              "a shortcut from the publisher name. For classical models, TF-IDF [12] vectors of "
              "unigrams and bigrams are extracted. For deep learning models, the text is converted "
              "to integer sequences and padded to a fixed length."),
        ("sh", "C. Classical Machine Learning Models"),
        ("p", "Logistic Regression, Multinomial Naive Bayes, Linear Support Vector Machine and "
              "Random Forest classifiers are trained on TF-IDF features. Hyper-parameters are "
              "selected using grid search with 5-fold cross-validation on the training set."),
        ("sh", "D. Hybrid CNN–LSTM Model"),
        ("p", "The input sequence is passed through an embedding layer initialised with "
              "pre-trained GloVe vectors. A one-dimensional convolution layer with max-pooling "
              "extracts local phrase features, and an LSTM layer then models the order of these "
              "features across the article. The output is passed through a dropout layer and a "
              "dense layer with sigmoid activation to predict the probability of the article "
              "being fake."),
        ("sh", "E. BERT-Based Model"),
        ("p", "The pre-trained bert-base-uncased model [10] is fine-tuned for binary "
              "classification. The article is tokenised with the WordPiece tokenizer and truncated "
              "to 512 tokens. The [CLS] representation is fed to a classification head, and the "
              "whole model is trained using the AdamW optimiser and binary cross-entropy loss."),
        ("sh", "F. Evaluation"),
        ("p", "All models are compared on the same test sets using accuracy, precision, recall, "
              "F1-score and the ROC-AUC. Since wrongly marking fake news as real is more harmful, "
              "recall of the fake class is given special importance."),
    ],
    "refs": [
        "K. Shu, A. Sliva, S. Wang, J. Tang, and H. Liu, “Fake news detection on social media: "
        "A data mining perspective,” ACM SIGKDD Explorations Newsletter, vol. 19, no. 1, "
        "pp. 22–36, 2017.",
        "S. Vosoughi, D. Roy, and S. Aral, “The spread of true and false news online,” "
        "Science, vol. 359, no. 6380, pp. 1146–1151, 2018.",
        "W. Y. Wang, “‘Liar, liar pants on fire’: A new benchmark dataset for fake "
        "news detection,” in Proc. ACL, 2017, pp. 422–426.",
        "V. Pérez-Rosas, B. Kleinberg, A. Lefevre, and R. Mihalcea, “Automatic detection "
        "of fake news,” in Proc. COLING, 2018, pp. 3391–3401.",
        "H. Ahmed, I. Traore, and S. Saad, “Detection of online fake news using n-gram "
        "analysis and machine learning techniques,” in Proc. ISDDC, LNCS vol. 10618, "
        "Springer, 2017, pp. 127–138.",
        "K. Shu, D. Mahudeswaran, S. Wang, D. Lee, and H. Liu, “FakeNewsNet: A data "
        "repository with news content, social context, and spatiotemporal information for "
        "studying fake news on social media,” Big Data, vol. 8, no. 3, pp. 171–188, 2020.",
        "N. Ruchansky, S. Seo, and Y. Liu, “CSI: A hybrid deep model for fake news "
        "detection,” in Proc. ACM CIKM, 2017, pp. 797–806.",
        "Y. Kim, “Convolutional neural networks for sentence classification,” in Proc. "
        "EMNLP, 2014, pp. 1746–1751.",
        "S. Hochreiter and J. Schmidhuber, “Long short-term memory,” Neural Computation, "
        "vol. 9, no. 8, pp. 1735–1780, 1997.",
        "J. Devlin, M.-W. Chang, K. Lee, and K. Toutanova, “BERT: Pre-training of deep "
        "bidirectional transformers for language understanding,” in Proc. NAACL-HLT, 2019, "
        "pp. 4171–4186.",
        "R. K. Kaliyar, A. Goswami, and P. Narang, “FakeBERT: Fake news detection in social "
        "media with a BERT-based deep learning approach,” Multimedia Tools and Applications, "
        "vol. 80, pp. 11765–11788, 2021.",
        "J. Ramos, “Using TF-IDF to determine word relevance in document queries,” in "
        "Proc. 1st Instructional Conf. Machine Learning, 2003.",
    ],
}

if __name__ == "__main__":
    build(PAPER1, "Paper1_Sentiment_Analysis_Rohit_Pujari.docx")
    build(PAPER2, "Paper2_Fake_News_Detection_Rohit_Pujari.docx")
