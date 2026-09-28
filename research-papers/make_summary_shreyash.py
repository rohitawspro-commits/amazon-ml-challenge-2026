"""Combined IEEE-format summary of two NLP papers (MT and summarization) for Shreyash Patil, Roll No. 18."""
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from make_papers import bullet, heading, para, set_cols, subheading

AUTHOR = ("Shreyash Patil", ["Roll No. 18, B.E. (Final Year)", "Department of Computer Engineering",
                             "[College Name], [City], India", "Guide: Prof. [Guide Name]"])
TITLE = ("A Summary of Two NLP Research Papers: English–Hindi Neural Machine Translation "
         "and Abstractive Text Summarization")

REFS = {
    "moses": "P. Koehn et al., “Moses: Open source toolkit for statistical machine translation,” in Proc. ACL "
             "Companion Volume (Demo and Poster Sessions), 2007, pp. 177–180.",
    "sutskever": "I. Sutskever, O. Vinyals, and Q. V. Le, “Sequence to sequence learning with neural networks,” "
                 "in Proc. NeurIPS, 2014, pp. 3104–3112.",
    "bahdanau": "D. Bahdanau, K. Cho, and Y. Bengio, “Neural machine translation by jointly learning to align and "
                "translate,” in Proc. ICLR, 2015.",
    "vaswani": "A. Vaswani et al., “Attention is all you need,” in Proc. NeurIPS, 2017, pp. 5998–6008.",
    "bpe": "R. Sennrich, B. Haddow, and A. Birch, “Neural machine translation of rare words with subword units,” "
           "in Proc. ACL, 2016, pp. 1715–1725.",
    "sentencepiece": "T. Kudo and J. Richardson, “SentencePiece: A simple and language independent subword tokenizer "
                     "and detokenizer for neural text processing,” in Proc. EMNLP: System Demonstrations, 2018, "
                     "pp. 66–71.",
    "backtrans": "R. Sennrich, B. Haddow, and A. Birch, “Improving neural machine translation models with monolingual "
                 "data,” in Proc. ACL, 2016, pp. 86–96.",
    "iitb": "A. Kunchukuttan, P. Mehta, and P. Bhattacharyya, “The IIT Bombay English-Hindi parallel corpus,” "
            "in Proc. LREC, 2018.",
    "samanantar": "G. Ramesh et al., “Samanantar: The largest publicly available parallel corpora collection for 11 "
                  "Indic languages,” Trans. ACL, vol. 10, pp. 145–162, 2022.",
    "mbart": "Y. Liu et al., “Multilingual denoising pre-training for neural machine translation,” Trans. ACL, "
             "vol. 8, pp. 726–742, 2020.",
    "bleu": "K. Papineni, S. Roukos, T. Ward, and W.-J. Zhu, “BLEU: A method for automatic evaluation of machine "
            "translation,” in Proc. ACL, 2002, pp. 311–318.",
    "sacrebleu": "M. Post, “A call for clarity in reporting BLEU scores,” in Proc. WMT, 2018, pp. 186–191.",
    "chrf": "M. Popović, “chrF: Character n-gram F-score for automatic MT evaluation,” in Proc. WMT, 2015, "
            "pp. 392–395.",
    "textrank": "R. Mihalcea and P. Tarau, “TextRank: Bringing order into text,” in Proc. EMNLP, 2004, "
                "pp. 404–411.",
    "nallapati": "R. Nallapati, B. Zhou, C. dos Santos, Ç. Gulçehre, and B. Xiang, “Abstractive text "
                 "summarization using sequence-to-sequence RNNs and beyond,” in Proc. CoNLL, 2016, pp. 280–290.",
    "cnndm": "K. M. Hermann et al., “Teaching machines to read and comprehend,” in Proc. NeurIPS, 2015, "
             "pp. 1693–1701.",
    "pointer": "A. See, P. J. Liu, and C. D. Manning, “Get to the point: Summarization with pointer-generator "
               "networks,” in Proc. ACL, 2017, pp. 1073–1083.",
    "bertsum": "Y. Liu and M. Lapata, “Text summarization with pretrained encoders,” in Proc. EMNLP-IJCNLP, "
               "2019, pp. 3730–3740.",
    "bart": "M. Lewis et al., “BART: Denoising sequence-to-sequence pre-training for natural language generation, "
            "translation, and comprehension,” in Proc. ACL, 2020, pp. 7871–7880.",
    "t5": "C. Raffel et al., “Exploring the limits of transfer learning with a unified text-to-text transformer,” "
          "J. Mach. Learn. Res., vol. 21, no. 140, pp. 1–67, 2020.",
    "pegasus": "J. Zhang, Y. Zhao, M. Saleh, and P. J. Liu, “PEGASUS: Pre-training with extracted gap-sentences for "
               "abstractive summarization,” in Proc. ICML, 2020, pp. 11328–11339.",
    "xsum": "S. Narayan, S. B. Cohen, and M. Lapata, “Don’t give me the details, just the summary! Topic-aware "
            "convolutional neural networks for extreme summarization,” in Proc. EMNLP, 2018, pp. 1797–1807.",
    "faithful": "J. Maynez, S. Narayan, B. Bohnet, and R. McDonald, “On faithfulness and factuality in abstractive "
                "summarization,” in Proc. ACL, 2020, pp. 1906–1919.",
    "rouge": "C.-Y. Lin, “ROUGE: A package for automatic evaluation of summaries,” in Proc. Text Summarization "
             "Branches Out (ACL Workshop), 2004, pp. 74–81.",
    "bertscore": "T. Zhang, V. Kishore, F. Wu, K. Q. Weinberger, and Y. Artzi, “BERTScore: Evaluating text "
                 "generation with BERT,” in Proc. ICLR, 2020.",
}

order = []


def cite(text):
    """[@a,@b] -> IEEE numbers assigned in order of first appearance."""
    def repl(m):
        nums = []
        for k in re.findall(r"@(\w+)", m[0]):
            if k not in order:
                order.append(k)
            nums.append(order.index(k) + 1)
        nums = sorted(set(nums))
        out, i = [], 0
        while i < len(nums):
            j = i
            while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
                j += 1
            out.append(f"[{nums[i]}]–[{nums[j]}]" if j - i >= 2 else ", ".join(f"[{n}]" for n in nums[i:j + 1]))
            i = j + 1
        return ", ".join(out)
    return re.sub(r"\[@\w+(?:,\s*@\w+)*\]", repl, text)


ABSTRACT = (
    "This document presents a combined summary of two research papers on sequence-to-sequence text generation, an "
    "important area of Natural Language Processing (NLP). The first paper proposes an English-to-Hindi neural "
    "machine translation system that trains a Transformer encoder–decoder on the IIT Bombay English–Hindi "
    "parallel corpus with SentencePiece subword units, compares it with phrase-based statistical and LSTM–attention "
    "baselines, and evaluates it with BLEU, chrF and human judgement. The second paper proposes an abstractive "
    "summarization system for English news articles that fine-tunes the pre-trained BART and T5 models on the "
    "CNN/DailyMail dataset and compares them with Lead-3, TextRank and pointer-generator baselines using ROUGE, "
    "BERTScore and a human check of factual consistency. For each paper the problem, literature survey, methodology "
    "and contributions are summarised; the two works are then compared, the techniques common to both are explained, "
    "and challenges and future directions are discussed.")
KEYWORDS = ("Natural Language Processing, Machine Translation, Text Summarization, Sequence-to-Sequence, Transformer, "
            "BART, T5, BLEU, ROUGE")

BODY = [
    ("h", "I. Introduction"),
    ("p", "Natural Language Processing (NLP) enables computers to understand and produce human language. Many NLP "
          "tasks, such as sentiment analysis or spam detection, map a text to a label. Other tasks must generate a "
          "new piece of text from an input text; these are called sequence-to-sequence (seq2seq) tasks. Machine "
          "translation and text summarization are the two best-known examples, and both are important in India, "
          "where information is produced in large volumes and in many languages."),
    ("p", "This document summarises two research papers on seq2seq generation. The first, “English–Hindi "
          "Neural Machine Translation Using the Transformer Architecture”, translates English sentences into "
          "Hindi. The second, “Abstractive Text Summarization of News Articles Using Pre-trained "
          "Sequence-to-Sequence Models”, writes short summaries of long news articles. Both papers build on the "
          "encoder–decoder Transformer [@vaswani]."),
    ("p", "Section II summarises the first paper and Section III the second. Section IV compares the two works, "
          "Section V explains the techniques common to both, Section VI discusses challenges and future scope, and "
          "Section VII concludes."),

    ("h", "II. Paper 1: English–Hindi Neural Machine Translation"),
    ("sh", "A. Problem Statement and Objectives"),
    ("p", "A large part of India's population is more comfortable reading Hindi than English, while most technical, "
          "legal and educational content on the web is available mainly in English. Machine translation (MT) "
          "converts text from a source language to a target language automatically. Earlier rule-based and "
          "phrase-based statistical MT systems [@moses] needed heavy feature engineering and handled the different "
          "word order of Hindi (subject–object–verb) and English (subject–verb–object) poorly. Neural MT "
          "instead trains one neural network end to end on sentence pairs [@sutskever, @bahdanau]."),
    ("p", "The objectives of the paper are:"),
    ("b", "To prepare a clean English–Hindi parallel corpus with subword tokenization."),
    ("b", "To train a Transformer encoder–decoder model for English-to-Hindi translation."),
    ("b", "To compare it with a phrase-based statistical baseline, an LSTM–attention baseline and a fine-tuned "
          "pre-trained multilingual model."),
    ("b", "To evaluate translations with BLEU, chrF and human ratings of adequacy and fluency."),
    ("sh", "B. Literature Survey"),
    ("p", "Koehn et al. [@moses] released Moses, the standard toolkit for phrase-based statistical MT. Sutskever et al. "
          "[@sutskever] showed that an LSTM encoder can compress a source sentence into a vector from which an LSTM "
          "decoder generates the translation. Bahdanau et al. [@bahdanau] removed the fixed-vector bottleneck with an "
          "attention mechanism that lets the decoder look at all encoder states, which greatly improved long "
          "sentences. Vaswani et al. [@vaswani] replaced recurrence completely with multi-head self-attention in the "
          "Transformer, which trains in parallel and became the standard NMT architecture."),
    ("p", "Sennrich et al. [@bpe] handled rare words by splitting them into subword units using byte-pair encoding, "
          "and Kudo and Richardson [@sentencepiece] made subword tokenization language independent with SentencePiece, "
          "which works directly on raw Devanagari text. Back-translation [@backtrans] creates extra synthetic training "
          "pairs from monolingual target-language text. For English–Hindi, Kunchukuttan et al. [@iitb] released the "
          "IIT Bombay parallel corpus, and Ramesh et al. [@samanantar] released Samanantar, a large collection for 11 "
          "Indic languages. Liu et al. [@mbart] showed that multilingual denoising pre-training (mBART) improves "
          "translation, especially for low-resource pairs. Translation quality is usually measured with BLEU "
          "[@bleu], reported with sacreBLEU [@sacrebleu] for comparable scores, and with chrF [@chrf], which uses "
          "character n-grams and suits morphologically rich languages such as Hindi."),
    ("sh", "C. Proposed Methodology"),
    ("lead", "1) Dataset: "),
    ("p", "The IIT Bombay English–Hindi corpus [@iitb] with its official training, development and test splits is "
          "used; its training part contains about 1.5 million sentence pairs. Duplicate and empty pairs, pairs longer "
          "than 100 tokens and pairs whose length ratio is above 2.5 are removed, and Devanagari text is Unicode-"
          "normalised because the same character can be stored in more than one way."),
    ("lead", "2) Tokenization: "),
    ("p", "Separate SentencePiece [@sentencepiece] BPE models with 16,000 subword units are trained for English and "
          "Hindi, since the two languages use different scripts."),
    ("lead", "3) Baselines: "),
    ("p", "A phrase-based statistical system built with Moses [@moses] and a bidirectional LSTM encoder–decoder "
          "with attention [@bahdanau] are trained on the same data."),
    ("lead", "4) Transformer Model: "),
    ("p", "The main model is the Transformer-base configuration [@vaswani]: six encoder and six decoder layers, model "
          "size 512, eight attention heads, feed-forward size 2048, dropout 0.1 and label smoothing 0.1, trained with "
          "Adam and a warm-up learning-rate schedule. Translations are produced with beam search (beam size 5) and a "
          "length penalty. In addition, the pre-trained mBART model [@mbart] is fine-tuned on the same data to measure "
          "the benefit of multilingual pre-training."),
    ("lead", "5) Evaluation: "),
    ("p", "BLEU [@bleu] and chrF [@chrf] are computed with sacreBLEU [@sacrebleu] on the test set. For human evaluation, "
          "100 random test sentences are rated from 1 to 5 for adequacy (meaning preserved) and fluency (natural Hindi) "
          "by two native speakers."),
    ("sh", "D. Key Contributions"),
    ("b", "A cleaning and subword pipeline designed for Devanagari text."),
    ("b", "A controlled comparison of statistical, recurrent, Transformer and pre-trained multilingual models on the "
          "same English–Hindi data."),
    ("b", "Evaluation that combines BLEU, the character-level chrF metric and human ratings."),

    ("h", "III. Paper 2: Abstractive Summarization of News Articles"),
    ("sh", "A. Problem Statement and Objectives"),
    ("p", "Thousands of news articles are published every day, and readers rarely have time to read them fully. "
          "Automatic text summarization produces a short version of a document that keeps its most important "
          "information. Extractive methods such as TextRank [@textrank] copy the most important sentences from the "
          "document, while abstractive methods write new sentences, as a human editor would, which gives shorter and "
          "more fluent summaries but is much harder."),
    ("p", "The objectives of the paper are:"),
    ("b", "To build an abstractive summarizer for English news articles."),
    ("b", "To compare extractive baselines, a pointer-generator network and fine-tuned pre-trained seq2seq Transformers."),
    ("b", "To evaluate summaries with ROUGE, BERTScore and a human check of factual consistency."),
    ("b", "To provide a web interface in which a user pastes an article and chooses a short or long summary."),
    ("sh", "B. Literature Survey"),
    ("p", "Mihalcea and Tarau [@textrank] proposed TextRank, which ranks sentences with a graph algorithm similar to "
          "PageRank. Nallapati et al. [@nallapati] applied attentional encoder–decoder RNNs to abstractive "
          "summarization and adapted the CNN/DailyMail question-answering corpus of Hermann et al. [@cnndm] into a "
          "multi-sentence summarization benchmark. See et al. [@pointer] proposed the pointer-generator network, which "
          "can copy words such as names and numbers from the article and uses a coverage mechanism to reduce "
          "repetition."),
    ("p", "Pre-trained Transformers changed the field. Liu and Lapata [@bertsum] used BERT-based encoders for both "
          "extractive and abstractive summarization. BART [@bart] is pre-trained as a denoising autoencoder that "
          "reconstructs corrupted text and performs strongly on summarization. T5 [@t5] treats every task as "
          "text-to-text, so summarization is requested with the prefix “summarize:”. PEGASUS [@pegasus] uses a "
          "pre-training objective designed for summarization that generates removed important sentences. Narayan et "
          "al. [@xsum] introduced XSum, which requires single-sentence summaries. Maynez et al. [@faithful] showed that "
          "abstractive models often “hallucinate” facts that are not in the source. Summaries are usually scored "
          "with ROUGE [@rouge], and BERTScore [@bertscore] compares texts using contextual embeddings instead of exact "
          "word overlap."),
    ("sh", "C. Proposed Methodology"),
    ("lead", "1) Dataset: "),
    ("p", "The non-anonymised CNN/DailyMail dataset [@cnndm, @nallapati] is used, with about 287,000 training, 13,000 "
          "validation and 11,000 test article–summary pairs. XSum [@xsum] is used as a second test set with very "
          "short summaries."),
    ("lead", "2) Preprocessing: "),
    ("p", "Bylines, source tags such as “(CNN) --” and HTML remains are removed. Articles are truncated to "
          "1024 tokens for BART and 512 tokens for T5, and reference summaries to 128 tokens."),
    ("lead", "3) Baselines: "),
    ("p", "Lead-3 (the first three sentences, a strong baseline for news), TextRank [@textrank] and a pointer-generator "
          "network with coverage [@pointer]."),
    ("lead", "4) Transformer Models: "),
    ("p", "BART [@bart] and T5-base [@t5] are fine-tuned with the AdamW optimiser, a small learning rate, label smoothing "
          "and early stopping on validation ROUGE-L. Summaries are generated with beam search (beam size 4), a length "
          "penalty, minimum and maximum length settings for short or long summaries, and blocking of repeated "
          "trigrams to avoid repetition."),
    ("lead", "5) Evaluation: "),
    ("p", "ROUGE-1, ROUGE-2 and ROUGE-L [@rouge] and BERTScore [@bertscore] are reported on the test sets. In addition, "
          "50 generated summaries are checked by hand, and every statement not supported by the article is marked as a "
          "hallucination, following Maynez et al. [@faithful]."),
    ("sh", "D. Key Contributions"),
    ("b", "A fair comparison of extractive, pointer-generator and pre-trained Transformer summarizers on the same data."),
    ("b", "Controllable summary length through decoding settings."),
    ("b", "An evaluation that measures factual consistency, not only word overlap."),
]

TABLE = [
    ("Aspect", "Paper 1: Machine Translation", "Paper 2: Summarization"),
    ("Task", "English sentence → Hindi sentence", "Long news article → short summary"),
    ("Input vs output", "About the same length", "Output much shorter than input"),
    ("Datasets", "IIT Bombay English–Hindi corpus", "CNN/DailyMail, XSum"),
    ("Baselines", "Moses phrase-based SMT, LSTM + attention", "Lead-3, TextRank, pointer-generator"),
    ("Main model", "Transformer-base; fine-tuned mBART", "Fine-tuned BART and T5"),
    ("Tokenization", "SentencePiece, separate vocabularies", "Vocabulary of the pre-trained model"),
    ("Metrics", "BLEU, chrF, human adequacy/fluency", "ROUGE, BERTScore, human factuality"),
    ("Main challenge", "Word order, rich Hindi morphology", "Hallucination, long inputs, repetition"),
    ("Application", "Content in Hindi for more readers", "Quick reading of news"),
]

AFTER = [
    ("p", "Table I shows that both papers solve a seq2seq problem with the same family of models: an encoder reads the "
          "input, a decoder generates the output token by token while attending to the encoder, and beam search "
          "chooses the final text. Both also compare classical baselines with Transformers and with pre-trained models."),
    ("p", "The goals of the two tasks are, however, opposite in one respect. A translation must keep all the meaning of "
          "the source, so leaving out information is an error. A summary must leave out most of the source and keep "
          "only the key facts. Copying also differs: in summarization, names and numbers are copied unchanged from the "
          "article [@pointer], while in translation names must be transliterated into Devanagari. Finally, both papers "
          "note that n-gram overlap metrics such as BLEU and ROUGE are limited, and therefore add chrF, BERTScore and "
          "human evaluation."),
    ("h", "V. Common NLP Techniques Used"),
    ("sh", "A. Subword Tokenization"),
    ("p", "Both systems split words into frequent subword units with BPE [@bpe] or SentencePiece [@sentencepiece]. "
          "Rare and unseen words, such as names or Hindi word forms, can then be built from known pieces, and the "
          "vocabulary stays small."),
    ("sh", "B. Encoder–Decoder with Attention"),
    ("p", "The encoder turns the input into a sequence of vectors, and the decoder generates the output one token at a "
          "time [@sutskever]. Attention [@bahdanau] lets every decoding step focus on the most relevant input "
          "positions. During training the decoder receives the correct previous token (teacher forcing) and is "
          "trained with cross-entropy loss, often with label smoothing."),
    ("sh", "C. The Transformer"),
    ("p", "The Transformer [@vaswani] uses self-attention, multi-head attention and positional encodings instead of "
          "recurrence. Every token can attend to every other token directly and training runs in parallel, which is "
          "why both papers use it as their main architecture."),
    ("sh", "D. Pre-training and Fine-tuning"),
    ("p", "Models such as mBART [@mbart], BART [@bart] and T5 [@t5] are first pre-trained on very large amounts of "
          "unlabelled text and then fine-tuned on the task data. This transfers general language knowledge and "
          "reduces the amount of labelled data needed."),
    ("sh", "E. Decoding"),
    ("p", "Greedy decoding picks the most probable token at each step, while beam search keeps several partial outputs "
          "and returns the best complete one. A length penalty controls output length, and n-gram blocking prevents "
          "the same phrase from being generated twice."),
    ("h", "VI. Challenges and Future Scope"),
    ("p", "The following challenges are common to both works:"),
    ("b", "Indian languages other than Hindi, and code-mixed Hinglish text, have far less training data."),
    ("b", "Generated text can be fluent but wrong: mistranslations and hallucinated facts [@faithful] are hard to detect."),
    ("b", "Self-attention cost grows with the square of input length, which limits very long documents."),
    ("b", "Automatic metrics do not fully agree with human judgement, so human evaluation remains necessary."),
    ("b", "Large models need GPUs; distillation and smaller models are needed for use on phones."),
    ("p", "Future work includes translation for more Indic languages using Samanantar [@samanantar], abstractive "
          "summarization of Hindi news, and cross-lingual summarization that reads an English article and writes the "
          "summary directly in Hindi, which combines the ideas of both papers. Summarization-specific pre-training "
          "such as PEGASUS [@pegasus] and back-translation [@backtrans] for low-resource pairs are also promising."),
    ("h", "VII. Conclusion"),
    ("p", "This document summarised two NLP research papers on English–Hindi neural machine translation and on "
          "abstractive summarization of news articles. Although one task preserves all the information and the other "
          "compresses it, both are solved with encoder–decoder Transformers, subword tokenization, attention, "
          "pre-training and beam search. The study shows that fine-tuned pre-trained seq2seq models are the most "
          "promising approach for both tasks, while careful evaluation with suitable metrics and human judgement is "
          "essential because generated text can be fluent but incorrect."),
]


def render(doc, items):
    lead = None
    for kind, text in items:
        if kind == "h":
            heading(doc, text)
        elif kind == "sh":
            subheading(doc, text)
        elif kind == "lead":
            lead = text
        elif kind == "b":
            bullet(doc, cite(text))
        else:
            p = para(doc)
            if lead:
                r = p.add_run(lead)
                r.font.size, r.italic = Pt(10), True
                lead = None
            r = p.add_run(cite(text))
            r.font.size = Pt(10)


def table(doc):
    para(doc, "Table I", size=8, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, indent=False, small_caps=True)
    para(doc, "Comparison of the Two Research Papers", size=8, align=WD_ALIGN_PARAGRAPH.CENTER, after=3,
         indent=False, small_caps=True)
    t = doc.add_table(rows=len(TABLE), cols=3)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = (Inches(0.75), Inches(1.25), Inches(1.25))
    for j, w in enumerate(widths):
        t.columns[j].width = w
    for i, row in enumerate(TABLE):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.width = widths[j]
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.size, r.bold = Pt(7.5), i == 0 or j == 0
    para(doc, "", after=4, indent=False)


def build(out):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.font.size = Pt(10)
    s = doc.sections[0]
    s.page_height, s.page_width = Inches(11.69), Inches(8.27)
    s.top_margin, s.bottom_margin = Inches(0.75), Inches(1.0)
    s.left_margin = s.right_margin = Inches(0.625)
    set_cols(s, 1)
    para(doc, TITLE, size=22, align=WD_ALIGN_PARAGRAPH.CENTER, after=12, indent=False)
    para(doc, AUTHOR[0], size=11, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    for ln in AUTHOR[1]:
        para(doc, ln, size=10, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    para(doc, "", after=6, indent=False)
    set_cols(doc.add_section(WD_SECTION.CONTINUOUS), 2, 360)
    for lead, body in (("Abstract—", ABSTRACT), ("Keywords—", KEYWORDS)):
        p = para(doc)
        r = p.add_run(lead)
        r.font.size, r.bold, r.italic = Pt(9), True, True
        r = p.add_run(body)
        r.font.size, r.bold = Pt(9), True
    render(doc, BODY)
    heading(doc, "IV. Comparative Analysis")
    table(doc)
    render(doc, AFTER)
    heading(doc, "References")
    for i, k in enumerate(order, 1):
        p = para(doc, indent=False)
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        r = p.add_run(f"[{i}]\t{REFS[k]}")
        r.font.size = Pt(8)
    unused = set(REFS) - set(order)
    assert not unused, f"uncited references: {unused}"
    doc.save(out)


if __name__ == "__main__":
    build("NLP_Research_Papers_Summary_Shreyash_Patil.docx")
