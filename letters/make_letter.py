"""Letter to the Principal proposing an in-house product development centre."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.font.size = Pt(11.5)
st.paragraph_format.space_after = Pt(0)
st.paragraph_format.line_spacing = 1.0
s = doc.sections[0]
s.page_height, s.page_width = Inches(11.69), Inches(8.27)
s.top_margin = s.bottom_margin = Inches(0.6)
s.left_margin = s.right_margin = Inches(0.9)


def para(text="", bold=False, after=0, align=WD_ALIGN_PARAGRAPH.LEFT, indent=None):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    if indent is not None:
        p.paragraph_format.left_indent = Inches(indent)
        p.paragraph_format.first_line_indent = Inches(-0.22)
    for k, part in enumerate(text.split("**")):
        if part:
            p.add_run(part).bold = bold or k % 2 == 1
    return p


J = WD_ALIGN_PARAGRAPH.JUSTIFY
for line in ["From,", "Rohit Pujari", "B.E. (Final Year), Department of Computer Engineering",
             "Roll No. 28", "[College Name], [City]"]:
    para(line)
para("Date: ____________", after=8)
for line in ["To,", "The Principal,", "[College Name], [City]"]:
    para(line)
para("Through: The Head of Department, Computer Engineering", after=8)
para("**Subject: Proposal to start an in-house Product Development Centre to provide internships and jobs "
     "to our own students**", after=8)
para("Respected Ma'am,", after=6)
para("I, Rohit Pujari, a final year student of the Computer Engineering department, would like to put forward a "
     "suggestion that I believe can solve the internship problem of our students and also create job "
     "opportunities within our own college.", align=J, after=6)
para("Medical colleges have their own attached hospital. Their students do their internship and practical "
     "training in the college's own hospital under the guidance of their teachers, so they never have to depend "
     "on outside hospitals for training. In engineering we do not have such a system, and every year many of "
     "our students find it difficult to get a good internship or a first job.", align=J, after=6)
para("In the same way, our college can start an **in-house Product Development Centre** that works like a small "
     "software company on campus. It can build real products, such as the college's own attendance, library or "
     "event management systems, and later websites and apps for local businesses and organisations. The centre "
     "can hire our brilliant students and fresh graduates who are genuinely interested in building products. "
     "Students who get offers from good outside companies should be free to join them; the centre only needs "
     "those who want to work on its products.", align=J, after=6)
para("**Benefits to the college and students:**", after=2)
for b in ["Every student can get a real, industry-level internship inside the college itself.",
          "Students gain practical experience, which improves their chances in placements.",
          "Talented students get a job opportunity immediately after their degree.",
          "The products and services can earn revenue for the college and help improve its labs.",
          "The college's reputation and its industry-interaction activities will grow.",
          "Senior students and graduates can guide juniors, so the work continues every year."]:
    para("•  " + b, align=J, indent=0.45)
para("", after=4)
para("To begin, the centre can be started on a small scale as a pilot, with 5 to 10 students under faculty "
     "mentors in one laboratory, and its first product can be a software solution needed by our own college. It "
     "can also work together with the college's innovation or incubation cell, if one exists.", align=J, after=6)
para("I humbly request you to consider this proposal and give me an opportunity to present a detailed plan "
     "at your convenience.", align=J, after=10)
for line in ["Thanking you,", "Yours faithfully,", "", "(Signature)", "Rohit Pujari",
             "B.E. Computer Engineering (Final Year), Roll No. 28"]:
    para(line)
doc.save("Letter_to_Principal_Rohit_Pujari.docx")
